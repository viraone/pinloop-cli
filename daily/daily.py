"""One command a day: pull manual-tester postings, score them, judge the best,
tailor a resume to each, and update the spreadsheet.

    daily/run.sh              the daily run
    daily/run.sh --no-pull    rebuild scores, resumes and the sheet for postings already held
    daily/run.sh --from "job boards"   pull from job boards instead of the automatic choice
"""
import argparse
import hashlib
import json
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import skills
import tailor

HERE = Path(__file__).parent
REPO = HERE.parent
PINLOOP = ["node", str(REPO / "dist" / "cli" / "pinloop.js")]
STATE = HERE / "state.json"
OUT = Path.home() / "Desktop" / "pinloop-jobs"
SHEET = Path.home() / "Desktop" / "pinloop-jobs.xlsx"
INDEX = OUT / "jobs.json"

TITLE = ("manual QA OR manual test OR manual tester OR manual testing OR QA analyst OR QA tester "
         "OR QA specialist OR software tester OR test analyst OR quality assurance analyst "
         "OR quality assurance tester OR quality analyst")
# Pulled first each day; only when these run dry does the broader TITLE fill the rest.
PRIORITY_TITLE = "manual tester OR quality assurance tester OR QA tester OR software tester OR manual QA tester"
FILTERS = ["--workplace", "Remote Solely", "--country", "United States"]
DAILY_LIMIT = 5
WINDOW_DAYS = 30
JUDGE_TOP = 3
SOURCES = ["career sites", "job boards"]

HDR = ["Added", "Company", "Title", "Location", "Posted", "Source", "ATS base", "ATS tailored",
       "Verdict", "Reasoning", "Tailored resume", "Apply link", "Status", "Pinloop ID"]
WIDTHS = [11, 20, 38, 20, 11, 12, 9, 9, 9, 70, 34, 55, 12, 38]


def log(msg: str) -> None:
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def run(*args: str, json_out: bool = True) -> dict | str:
    """Runs pinloop; JSON from stdout, anything for a person from stderr goes to the log."""
    p = subprocess.run(PINLOOP + list(args) + (["--json"] if json_out else ["--plain"]),
                       capture_output=True, text=True)
    for line in p.stderr.splitlines():
        if line.strip():
            log(f"  pinloop: {line.strip()}")
    if p.returncode != 0:
        raise SystemExit(f"pinloop {' '.join(args[:2])} failed:\n{p.stdout}\n{p.stderr}")
    if not json_out:
        return p.stdout
    try:
        return json.loads(p.stdout)
    except json.JSONDecodeError:
        raise SystemExit(f"pinloop {args[0]} did not print JSON:\n{p.stdout}\n{p.stderr}")


def postings_left() -> int:
    text = run(json_out=False)
    m = re.search(r"Postings: \d+ used, (\d+) left", text)
    for line in text.splitlines():
        if line.startswith("Tell the person"):
            log(line)
    return int(m.group(1)) if m else 0


def load_state() -> dict:
    if STATE.exists():
        return json.loads(STATE.read_text())
    return {"last_from": SOURCES[1]}


def held_ids() -> set[str]:
    return {r["id"] for r in run("viewed", "--all").get("rows", [])}


def pull(state: dict, source: str, limit: int, priority: bool = False, held: set[str] | None = None) -> list[dict]:
    """Newest matching postings this account has never been handed, via search.

    A search charges only rows new to the account (a pull charges every row, held or
    not), so pages are walked until `limit` new rows are in hand or the matches run out."""
    since = (date.today() - timedelta(days=WINDOW_DAYS)).isoformat()
    title = PRIORITY_TITLE if priority else TITLE
    log(f"searching {source} for the newest {'tester-titled' if priority else 'broader'} postings since {since}")
    held = set(held or ())
    new_rows: list[dict] = []
    cursor = None
    for _ in range(4):
        want = limit - len(new_rows)
        if want <= 0:
            break
        args = ["search", "--in", "title", title, *FILTERS, "--from", source, "--posted-after", since,
                "--order", "newest", "--limit", str(want)]
        if cursor:
            args += ["--cursor", cursor]
        answer = run(*args)
        rows = answer.get("rows", [])
        if answer.get("refused"):
            log(f"  refused: {answer['refused'].splitlines()[0]}")
            break
        fresh = []
        for r in rows:
            if r["id"] in held or r.get("already_had"):
                continue
            fixture = not_hiring(r)
            if fixture:
                log(f"  not a real opening, skipped: {r.get('company')} — {r.get('title')} ({fixture})")
                continue
            fresh.append(r)
        for r in fresh:
            r["_source"] = source
            held.add(r["id"])
        new_rows += fresh
        log(f"  page: {len(rows)} rows, {len(fresh)} new")
        cursor = answer.get("cursor")
        if not rows or not cursor:
            break
    state["last_from"] = source
    return new_rows


def ats(ids: list[str], resume_file: Path | None = None) -> dict[str, dict]:
    if not ids:
        return {}
    extra = ["--resume", str(resume_file)] if resume_file else []
    answer = run("ats", *ids, *extra)
    return {r["id"]: r for r in answer["rows"]}


def background_text() -> str:
    """What this person has done, for the judge: the resume's facts plus their own proof sentences."""
    r = tailor.base_resume()
    out = [r["summary"], ""]
    for job in r["experience"]:
        out.append(f"{job['title']} at {job['company']} ({job['dates']}): " + " ".join(job["bullets"]))
    out.append("")
    out.append("Skills, every one backed by the work above or confirmed by the person: "
               + "; ".join(f"{g}: {', '.join(i)}" for g, i in r["skills"]))
    if r.get("statements"):
        out += ["", "In the person's own words:"] + r["statements"]
    out += ["", "Wants: fully remote (US), manual / mobile QA, test analyst or QA analyst roles. "
            "Does not want automation-first, SDET, security or red-team roles, and will not claim coding or "
            "automation-framework experience."]
    return "\n".join(out)


def sync_profile_resume(pdf: Path, state: dict) -> None:
    """Pinloop's judge reads the stored resume, so it gets the same base resume the scores use."""
    digest = hashlib.sha256((pdf.with_suffix(".txt")).read_bytes()).hexdigest()
    if state.get("resume_sha") == digest and state.get("background_v") == 1:
        return
    run("profile", "put", "resume", "--file", str(pdf), json_out=False)
    bg = HERE / "background.txt"
    bg.write_text(background_text())
    run("profile", "put", "background", "--file", str(bg), json_out=False)
    state["resume_sha"] = digest
    state["background_v"] = 1
    STATE.write_text(json.dumps(state, indent=1))
    log("base resume changed; stored the new one on Pinloop for judging")


def stored_verdicts() -> dict[str, dict]:
    return {r["id"]: r for r in run("judgment", "list").get("rows", [])}


def parse_reasoning(text: str) -> dict:
    """Callback odds and the three fixes out of a recruiter-eye reasoning; empty when it is the old free text."""
    m = re.search(r"callback odds:\s*(\d{1,3})\s*%", text, re.I)
    fixes = [f.strip() for f in re.findall(r"Fix\s*\d\s*:\s*(.+?)(?=\n\s*Fix\s*\d\s*:|\Z)", text, re.I | re.S)]
    body = re.sub(r"callback odds:.*?\n", "", text, count=1, flags=re.I)
    body = re.split(r"\n\s*Fix\s*1\s*:", body, maxsplit=1, flags=re.I)[0].strip()
    return {"odds": int(m.group(1)) if m else None, "fixes": fixes[:3], "summary": body if m else ""}


def judge(ids: list[str], again: bool = False) -> None:
    if not ids:
        return
    extra = ["--again"] if again else []
    first = run("judge", *ids, *extra)
    token = first.get("confirm_token")
    if not token:
        return
    left = first.get("confirm_left", {}).get("full", "?")
    log(f"judging {len(ids)} postings ({left} full judgments left this month before this)")
    run("judge", *ids, *extra, "--confirm", token)


def judge_targets(ids: list[str], by_id: dict[str, dict], scores: dict[str, dict],
                  have: dict[str, dict], top: int) -> tuple[list[str], dict[str, str]]:
    """The postings worth paying a judgment for, and why each of the rest goes without.

    A judgment buys a verdict for a role, not for a URL, so a role posted twice must not buy
    it twice: one best-scoring stand-in per role competes for the day's slots. A row goes
    without because its stand-in was judged instead ("duplicate") or because the role itself
    lost the slice ("cut") -- never assumed, so the app can say which it was."""
    want = [i for i in sorted(ids, key=lambda i: -scores[i]["ats_score"]) if i not in have]
    stand_in: dict[tuple[str, str], str] = {}
    for i in want:
        stand_in.setdefault(role_key(by_id[i]), i)
    picked = list(stand_in.values())[:top]
    why = {}
    for i in want:
        if i not in picked:
            why[i] = "duplicate" if stand_in[role_key(by_id[i])] in picked else "cut"
    return picked, why


def folder_name(s: str) -> str:
    """A folder name a person reads: the company or title itself, minus what a filesystem refuses."""
    return re.sub(r"[\\/:*?\"<>|]+", "-", s).strip(" .")[:80] or "Untitled"


# Boards that exist to exercise an integration rather than to hire. Their postings score and
# tailor like real ones, which is the trouble: the resume work looks done and Apply leads nowhere.
# Judged only on the board token an applicant tracker puts in the URL, where names are concatenated
# and the signal is strong. Company names are left alone on purpose: matching those caught nothing
# real and cost false positives, since Sandboxx and Testlio are employers and "test" is this
# feed's whole subject.
ATS_BOARD = re.compile(r"(?:greenhouse\.io|lever\.co|ashbyhq\.com|workable\.com"
                       r"|smartrecruiters\.com|myworkdayjobs\.com)/(?:embed/job_app\?for=)?"
                       r"([A-Za-z0-9_-]+)", re.I)
# "sandbox" must end the word: builtinintegrationsandbox is a fixture, Sandboxx is an employer.
FIXTURE = re.compile(r"sandbox(?![a-z0-9])|sandpit|integration-?test|\bdemo\b", re.I)


def not_hiring(row: dict) -> str:
    """Why a posting is a fixture rather than a job, or "" when it looks like a real opening."""
    for field in ("url", "posting_url"):
        board = ATS_BOARD.search(str(row.get(field) or ""))
        if board and FIXTURE.search(board.group(1)):
            return f"board {board.group(1).lower()!r}"
    return ""


def role_key(row: dict) -> tuple[str, str]:
    """Company and title folded together; two postings sharing it are one opening advertised twice."""
    company = str(row.get("company") or "").strip().casefold()
    title = tailor.clean_title(str(row.get("title") or "")).strip().casefold()
    return company, title


def posting_stamp(row: dict) -> str:
    """Six stable characters telling apart two postings a company gave the same title.

    Taken from the posting itself, never from the score: a keyword fingerprint would move
    every time the base resume changes, and every stored resume path would go stale with it."""
    seed = str(row.get("url") or row.get("posting_url") or row.get("id") or "")
    return hashlib.sha256(seed.encode()).hexdigest()[:6]


def load_index() -> dict[str, dict]:
    if INDEX.exists():
        return {j["id"]: j for j in json.loads(INDEX.read_text())}
    return {}


def save_index(jobs: dict[str, dict]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    order = {"strong": 0, "fair": 1, "weak": 2, "no": 3}
    ordered = sorted(jobs.values(), key=lambda j: (order.get(j["verdict"], 4), -(j["ats_tailored"] or 0), j["added"]))
    INDEX.write_text(json.dumps(ordered, indent=1))


def load_sheet() -> tuple[Workbook, dict[str, list]]:
    """Existing rows keyed by Pinloop ID; a sheet with another layout is carried over by id."""
    existing: dict[str, list] = {}
    if SHEET.exists():
        old = load_workbook(SHEET).active
        header = [c.value for c in old[1]]
        for row in old.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(header, row))
            pid = rec.get("Pinloop ID")
            if pid:
                existing[pid] = [rec.get(h) for h in HDR]
                if header != HDR:
                    existing[pid][HDR.index("Status")] = rec.get("Status")
                    existing[pid][HDR.index("Apply link")] = rec.get("URL") or rec.get("Apply link")
    wb = Workbook()
    ws = wb.active
    ws.title = "Postings"
    ws.append(HDR)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A2B30")
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return wb, existing


def save_sheet(wb: Workbook, rows: dict[str, list]) -> None:
    ws = wb.active
    order = {"strong": 0, "fair": 1, "weak": 2, "no": 3}
    for r in sorted(rows.values(), key=lambda r: (order.get(r[8], 4), -(r[7] or 0), str(r[0]))):
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws["L"][1:]:
        if c.value:
            c.hyperlink = c.value
            c.font = Font(color="0563C1", underline="single")
    wb.save(SHEET)


def notify(text: str) -> None:
    subprocess.run(["osascript", "-e",
                    f'display notification "{text}" with title "Pinloop daily"'], check=False)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pull", action="store_true", help="rebuild for postings already held, pull nothing")
    ap.add_argument("--from", dest="source", choices=SOURCES, help="pull from this source instead of alternating")
    ap.add_argument("--judge-top", type=int, default=JUDGE_TOP, help="judge this many of today's best by ATS score")
    ap.add_argument("--rejudge", action="store_true", help="buy fresh recruiter-eye verdicts for postings already judged")
    args = ap.parse_args()

    state = load_state()
    today = date.today().isoformat()
    new_rows: list[dict] = []

    if args.no_pull:
        held = run("viewed", "--all").get("rows", [])
        for r in held:
            r["_source"] = ""
        new_rows = held
        log(f"rebuilding for {len(held)} postings already held")
    else:
        left = postings_left()
        if left == 0:
            log("no postings left today; nothing pulled. Usage resets at midnight UTC (5pm Pacific).")
            notify("Nothing pulled: today's postings are used up.")
            return
        source = args.source or SOURCES[(SOURCES.index(state["last_from"]) + 1) % len(SOURCES)]
        other = SOURCES[(SOURCES.index(source) + 1) % len(SOURCES)]
        want = min(DAILY_LIMIT, left)
        held = held_ids()
        # Tester-titled postings first, from both places; the broader query only fills what is left.
        plan = [(source, True), (other, True), (source, False), (other, False)]
        if args.source:
            plan = [(source, True), (source, False)]
        for src, prio in plan:
            short = want - len(new_rows)
            if short <= 0:
                break
            got = pull(state, src, short, priority=prio, held=held)
            # pull() works off a copy, so what it just took has to be fed back in by hand;
            # without this a posting carried by both sources is pulled, scored and tailored twice.
            held.update(r["id"] for r in got)
            new_rows += got
        STATE.write_text(json.dumps(state, indent=1))

    real = [r for r in new_rows if not not_hiring(r)]
    if len(real) < len(new_rows):
        log(f"{len(new_rows) - len(real)} posting(s) from a test board left out")
        new_rows = real

    if not new_rows:
        log("nothing new today")
        notify("No new manual-tester postings today.")
        return

    ids = [r["id"] for r in new_rows]
    base_pdf, base_txt = tailor.build_base(OUT / "_base resume")
    sync_profile_resume(base_pdf, state)
    base_scores = ats(ids, base_txt)

    have = stored_verdicts()
    if args.rejudge:
        stale = [i for i in ids if i in have and parse_reasoning(have[i].get("reasoning", ""))["odds"] is None]
        judge(stale, again=True)
    # Why a posting has no recruiter's read, so the app can say which it is instead of guessing.
    why_unjudged: dict[str, str] = {}
    if not args.no_pull:
        by_id = {r["id"]: r for r in new_rows}
        picked, why_unjudged = judge_targets(ids, by_id, base_scores, have, args.judge_top)
        twins = sum(1 for w in why_unjudged.values() if w == "duplicate")
        if twins:
            log(f"{twins} duplicate posting(s) of a role being judged; not paying for that verdict twice")
        judge(picked)
    verdicts = stored_verdicts()

    wb, rows = load_sheet()
    jobs = load_index()
    for pid, j in list(jobs.items()):
        fixture = not_hiring(j)
        if fixture:
            del jobs[pid]
            rows.pop(pid, None)
            log(f"dropped from the dashboard, not a real opening: {j['company']} — {j['title']} ({fixture})")
    for r in new_rows:
        pid = r["id"]
        rep = base_scores.get(pid)
        if not rep:
            continue
        company = str(r.get("company") or "")
        title = tailor.clean_title(str(r.get("title") or ""))
        folder = OUT / folder_name(company) / f"{folder_name(title)} ({posting_stamp(r)})"
        pdf, txt = tailor.build(r, rep, folder, tailor.resume_name())
        tailored = ats([pid], txt).get(pid, {}).get("ats_score")
        v = verdicts.get(pid, {})
        old = rows.get(pid)
        prior = jobs.get(pid, {})
        if v.get("verdict"):
            why = ""
        elif args.no_pull:
            # A row first seen in a rebuild genuinely never went up for judging. One carried over
            # from before this field existed is simply unknown, and says so rather than guessing.
            why = prior.get("why_unjudged") or ("" if prior else "rebuild")
        else:
            why = why_unjudged.get(pid, "")
        jobs[pid] = {
            "id": pid, "added": prior.get("added") or today, "company": company, "title": title,
            "location": ", ".join(r.get("locations") or r.get("countries") or []),
            "posted": str(r.get("posted_at") or "")[:10], "source": r.get("_source") or prior.get("source", ""),
            "ats_base": rep["ats_score"], "ats_tailored": tailored,
            "verdict": v.get("verdict", ""), "reasoning": v.get("reasoning", ""), "why_unjudged": why,
            **parse_reasoning(v.get("reasoning", "")),
            "matched": [{"k": tailor.label_for(h), "r": h["resume"], "p": h["posting"], "kind": h["kind"]} for h in rep["ats_matched"]],
            "missing": [{"k": tailor.label_for(h), "p": h["posting"], "kind": h["kind"]} for h in rep["ats_missing"]],
            "resume": str(pdf), "folder": str(folder), "url": r.get("url") or r.get("posting_url") or "",
            "status": prior.get("status") or (old[12] if old and old[12] else "Not applied"),
        }
        rows[pid] = [
            old[0] if old and old[0] else today,
            company, title, jobs[pid]["location"], jobs[pid]["posted"], jobs[pid]["source"],
            rep["ats_score"], tailored, v.get("verdict", ""), v.get("reasoning", ""),
            str(pdf.relative_to(Path.home() / "Desktop")), jobs[pid]["url"], jobs[pid]["status"], pid,
        ]
        log(f"  {rep['ats_score']:>3} → {tailored:>3}  {v.get('verdict', '-'):6} {company} — {title}")
        top_missing = [h["posting_spelling"] for h in rep["ats_missing"] if h["kind"] != "soft"][:6]
        if top_missing:
            log(f"       missing from resume: {', '.join(top_missing)}")
    save_sheet(wb, rows)
    save_index(jobs)
    ask = skills.pending(list(jobs.values()), tailor.base_resume())
    if ask:
        log(f"skills check: {len(ask)} to confirm in the app — " + ", ".join(a["skill"] for a in ask[:8]))
    log(f"spreadsheet updated: {SHEET}")
    log(f"tailored resumes in: {OUT}/<Company>/<Job Title>/{tailor.resume_name()}.pdf")
    notify(f"{len(new_rows)} postings added. Spreadsheet and resumes are ready.")


if __name__ == "__main__":
    sys.exit(main())
