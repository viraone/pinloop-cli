"""The local web app: a Play button that runs the daily pipeline, and a card per job.

    daily/ui.sh          starts it and opens http://127.0.0.1:5177
"""
import json
import re
import urllib.request
import os
import subprocess
import sys
import threading
import webbrowser
from datetime import datetime
from pathlib import Path

from flask import Flask, abort, jsonify, request, send_file, send_from_directory
from openpyxl import load_workbook

import skills
import tailor

HERE = Path(__file__).parent
PY = HERE / ".venv" / "bin" / "python"
OUT = Path.home() / "Desktop" / "pinloop-jobs"
INDEX = OUT / "jobs.json"
SHEET = Path.home() / "Desktop" / "pinloop-jobs.xlsx"
PORT = 5177

app = Flask(__name__, static_folder=None)
state = {"running": False, "log": [], "started": None, "finished": None, "ok": None}
lock = threading.Lock()


def jobs() -> list[dict]:
    return json.loads(INDEX.read_text()) if INDEX.exists() else []


def save_jobs(all_jobs: list[dict]) -> None:
    INDEX.write_text(json.dumps(all_jobs, indent=1))


def pinloop_status() -> str:
    p = subprocess.run(["node", str(HERE.parent / "dist" / "cli" / "pinloop.js"), "--plain"],
                       capture_output=True, text=True)
    keep = [l for l in p.stdout.splitlines() if l.startswith(("Postings:", "Judging:"))]
    return "\n".join(keep)


def run_pipeline(args: list[str]) -> None:
    with lock:
        state.update(running=True, log=[], started=datetime.now().isoformat(), finished=None, ok=None)
    p = subprocess.Popen([str(PY), str(HERE / "daily.py"), *args], cwd=HERE,
                         stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    for line in p.stdout:
        with lock:
            state["log"].append(line.rstrip())
    p.wait()
    with lock:
        state.update(running=False, finished=datetime.now().isoformat(), ok=p.returncode == 0)


@app.get("/")
def home():
    return send_from_directory(HERE / "static", "index.html")


@app.get("/api/jobs")
def api_jobs():
    last = datetime.fromtimestamp(INDEX.stat().st_mtime).isoformat() if INDEX.exists() else None
    return jsonify({"jobs": jobs(), "status": pinloop_status(), "last_updated": last})


@app.post("/api/run")
def api_run():
    if state["running"]:
        return jsonify({"ok": False, "why": "already running"}), 409
    args = []
    body = request.get_json(silent=True) or {}
    if body.get("no_pull"):
        args.append("--no-pull")
    if body.get("source"):
        args += ["--from", body["source"]]
    threading.Thread(target=run_pipeline, args=(args,), daemon=True).start()
    return jsonify({"ok": True})


@app.get("/api/run/status")
def api_run_status():
    with lock:
        return jsonify(state)


def job_or_404(job_id: str) -> dict:
    for j in jobs():
        if j["id"] == job_id:
            return j
    abort(404)


@app.get("/resume/<job_id>")
def resume(job_id):
    j = job_or_404(job_id)
    return send_file(j["resume"], mimetype="application/pdf", download_name=Path(j["resume"]).name)


@app.post("/api/apply/<job_id>")
def api_apply(job_id):
    j = job_or_404(job_id)
    if not j.get("url"):
        return jsonify({"ok": False, "why": "no apply link"}), 400
    subprocess.Popen([str(PY), str(HERE / "apply.py"), j["url"], j["resume"]],
                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
    return jsonify({"ok": True})


@app.post("/api/reveal/<job_id>")
def api_reveal(job_id):
    j = job_or_404(job_id)
    subprocess.run(["open", "-R", j["resume"]], check=False)
    return jsonify({"ok": True})


@app.post("/api/status/<job_id>")
def api_status(job_id):
    new = (request.get_json(silent=True) or {}).get("status", "Not applied")
    all_jobs = jobs()
    for j in all_jobs:
        if j["id"] == job_id:
            j["status"] = new
    save_jobs(all_jobs)
    if SHEET.exists():
        wb = load_workbook(SHEET)
        ws = wb.active
        header = [c.value for c in ws[1]]
        if "Pinloop ID" in header and "Status" in header:
            id_col, st_col = header.index("Pinloop ID") + 1, header.index("Status") + 1
            for row in ws.iter_rows(min_row=2):
                if row[id_col - 1].value == job_id:
                    row[st_col - 1].value = new
        wb.save(SHEET)
    return jsonify({"ok": True})


@app.post("/api/open-sheet")
def api_open_sheet():
    subprocess.run(["open", str(SHEET)], check=False)
    return jsonify({"ok": True})


@app.get("/api/skills")
def api_skills():
    base = tailor.base_resume()
    ask = skills.pending(jobs(), json.loads((HERE / "resume.json").read_text()))
    for a in ask:
        a["question"] = skills.question(a["skill"])
    d = skills.load()
    added = [{"skill": s, "note": n} for s, n in skills.additional(json.loads((HERE / "resume.json").read_text()))]
    return jsonify({"pending": ask, "added": added,
                    "skipped": [{"skill": skills.pretty(s), "note": n} for s, n in d["not"].items()],
                    "base_skill_count": sum(len(i) for _, i in base["skills"])})


@app.get("/api/evidence")
def api_evidence():
    base = json.loads((HERE / "resume.json").read_text())
    return jsonify({"rows": skills.evidence(base)})


@app.post("/api/statement")
def api_statement():
    body = request.get_json(silent=True) or {}
    skill, text = body.get("skill", "").strip(), body.get("text", "")
    if not skill:
        abort(400)
    skills.set_statement(skill, text)
    d = skills.load()
    n = skills.norm(skill)
    if text.strip() and n not in d["have"]:
        d["not"].pop(n, None)
        d["have"][n] = "you wrote a proof sentence in the app"
        skills.save(d)
    return jsonify({"ok": True})


OLLAMA = "http://127.0.0.1:11434"
TIDY_MODELS = ["llama3.1:8b", "gpt-oss:20b", "llama3.3:latest"]
TIDY_PROMPT = """You rewrite a job seeker's spoken sentence into ONE resume bullet that passes ATS keyword scanners.
Rules:
- Keep every fact exactly as spoken. Never add tools, numbers, years, or companies the person did not say.
- Include the skill name "{skill}" verbatim.
- Start with a strong past-tense action verb (Tested, Verified, Documented, Analyzed, Coordinated...). No "I", "my", "we".
- Plain ATS-safe wording, no jargon fillers, no quotes, no emojis, 12-30 words, one sentence, no trailing period needed.
- If the sentence has no real experience in it (e.g. "I don't know"), reply with just the word: EMPTY
Reply with the bullet only."""


def _tidy_fallback(skill, text):
    t = " ".join(text.split()).strip(" .")
    t = re.sub(r"^(um+|uh+|so|like|well|okay|ok)[, ]+", "", t, flags=re.I)
    t = re.sub(r"^(i|i've|i have|i've been|i was|i am|i'm)\s+", "", t, flags=re.I)
    t = t[:1].upper() + t[1:]
    if skill.lower() not in t.lower():
        t = f"{t} using {skill}"
    return t


@app.post("/api/tidy")
def api_tidy():
    body = request.get_json(silent=True) or {}
    skill, text = body.get("skill", "").strip(), " ".join(body.get("text", "").split())
    if not skill or not text:
        abort(400)
    for model in TIDY_MODELS:
        try:
            r = urllib.request.Request(f"{OLLAMA}/api/chat", data=json.dumps({
                "model": model, "stream": False, "options": {"temperature": 0.2, "num_predict": 80},
                "messages": [{"role": "system", "content": TIDY_PROMPT.format(skill=skill)},
                             {"role": "user", "content": text}]}).encode(), headers={"content-type": "application/json"})
            with urllib.request.urlopen(r, timeout=60) as resp:
                out = json.loads(resp.read())["message"]["content"].strip().strip('"').strip("•-* ").strip()
            out = out.split("\n")[0].strip()
            if out.upper() == "EMPTY":
                return jsonify({"text": "", "empty": True, "model": model})
            if out:
                return jsonify({"text": out, "model": model})
        except Exception as e:  # model missing / ollama down → try next, then fallback
            if os.environ.get("PL_DEBUG"):
                print("tidy", model, e)
    return jsonify({"text": _tidy_fallback(skill, text), "model": "basic"})


@app.get("/api/profile")
def api_profile():
    return jsonify(json.loads((HERE / "profile.json").read_text()))


@app.post("/api/profile")
def api_profile_save():
    body = request.get_json(silent=True) or {}
    cur = json.loads((HERE / "profile.json").read_text())
    for k in cur:
        if k in body:
            cur[k] = str(body[k])
    (HERE / "profile.json").write_text(json.dumps(cur, indent=2, ensure_ascii=False) + "\n")
    return jsonify({"ok": True})


@app.post("/api/skills/<verdict>")
def api_skills_decide(verdict):
    if verdict not in ("have", "not", "undo"):
        abort(400)
    skill = (request.get_json(silent=True) or {}).get("skill", "").strip()
    if not skill:
        abort(400)
    d = skills.load()
    n = skills.norm(skill)
    d["have"].pop(n, None)
    d["not"].pop(n, None)
    if verdict != "undo":
        d[verdict][n] = "you confirmed this in the app" if verdict == "have" else "you said no in the app"
    skills.save(d)
    return jsonify({"ok": True})


if __name__ == "__main__":
    if "--no-browser" not in sys.argv:
        threading.Timer(0.8, lambda: webbrowser.open(f"http://127.0.0.1:{PORT}")).start()
    app.run(host="127.0.0.1", port=PORT, debug=False, threaded=True)
